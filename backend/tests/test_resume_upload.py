from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.core.database import SessionLocal
from app.main import app
from app.models.auth import User
from app.models.evaluation import Evaluation
from app.core.config import settings


def _sample_docx(name: str, body: str) -> tuple[str, bytes, str]:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(name)
    document.add_paragraph(body)
    document.save(buffer)
    return (
        'resume.docx',
        buffer.getvalue(),
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


def test_resume_upload_persists_candidate_with_mock_llm():
    client = TestClient(app)
    email = f'{uuid4().hex}@example.com'
    assert client.post('/api/auth/register', json={'email': email, 'password': 'correct horse battery'}).status_code == 201

    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        evaluation = Evaluation(
            user_id=user.id,
            job_title='Backend Engineer',
            job_description='Build APIs with Python and FastAPI.',
            structured_job={
                'role': 'Backend Engineer',
                'required_skills': ['Python', 'FastAPI'],
                'preferred_skills': [],
                'minimum_experience': 2,
                'education_requirements': [],
                'responsibilities': ['Build REST APIs'],
                'required_qualifications': [],
                'preferred_qualifications': [],
                'relevant_certifications': [],
                'seniority': None,
                'explicit_experience_requirements': None,
            },
            status='pending',
        )
        db.add(evaluation)
        db.commit()
        db.refresh(evaluation)
        evaluation_id = evaluation.id
    finally:
        db.close()

    file_name, file_bytes, mime = _sample_docx('Alex Rivera', 'Python and FastAPI developer with 5 years experience.')
    response = client.post(
        f'/api/evaluations/{evaluation_id}/resumes',
        files={'files': (file_name, file_bytes, mime)},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload['evaluation_id'] == evaluation_id
    assert len(payload['results']) == 1
    assert payload['results'][0]['status'] == 'completed', payload['results'][0].get('error')
    assert payload['results'][0]['score'] is not None

    detail = client.get(f'/api/evaluations/{evaluation_id}')
    assert detail.status_code == 200
    assert len(detail.json()['candidates']) == 1


def test_resume_upload_records_oversized_file_as_failed_candidate(monkeypatch):
    client = TestClient(app)
    email = f'{uuid4().hex}@example.com'
    assert client.post('/api/auth/register', json={'email': email, 'password': 'correct horse battery'}).status_code == 201

    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        evaluation = Evaluation(
            user_id=user.id,
            job_title='Backend Engineer',
            job_description='Build APIs with Python and FastAPI.',
            structured_job={
                'role': 'Backend Engineer',
                'required_skills': ['Python'],
                'preferred_skills': [],
                'minimum_experience': None,
                'education_requirements': [],
                'responsibilities': [],
                'required_qualifications': [],
                'preferred_qualifications': [],
                'relevant_certifications': [],
                'seniority': None,
                'explicit_experience_requirements': None,
            },
            status='pending',
        )
        db.add(evaluation)
        db.commit()
        db.refresh(evaluation)
        evaluation_id = evaluation.id
    finally:
        db.close()

    monkeypatch.setattr(settings, 'max_upload_size_mb', 1)
    response = client.post(
        f'/api/evaluations/{evaluation_id}/resumes',
        files={'files': ('large.pdf', b'%PDF' + b'x' * (1024 * 1024 + 1), 'application/pdf')},
    )

    assert response.status_code == 200
    assert response.json()['results'][0]['status'] == 'failed'
    assert 'upload limit' in response.json()['results'][0]['error']
    assert client.get(f'/api/evaluations/{evaluation_id}').json()['status'] == 'failed'
