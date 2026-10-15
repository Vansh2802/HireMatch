from io import BytesIO

from docx import Document

from app.services.document_parser import parse_uploaded_document


def test_docx_parsing_uses_in_memory_stream():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph('Jordan Lee')
    document.add_paragraph('Python engineer with FastAPI experience.')
    document.save(buffer)

    parsed = parse_uploaded_document('resume.docx', buffer.getvalue())

    assert 'Jordan Lee' in parsed['text']
    assert 'Python' in parsed['text']
